from docx import Document
from docx.shared import Pt, RGBColor

doc = Document()

# 标题
doc.add_heading('行业分析 · 智能家居售后维修赛道', level=1)

doc.add_heading('一、市场规模', level=2)
doc.add_paragraph('2025 年中国智能家居市场规模达 6500 亿元，年复合增长率 18%')
doc.add_paragraph('全国智能家居设备存量超 2.3 亿台，按 8% 年故障率计算，年维修需求约 1800 万单')
doc.add_paragraph('单均客单价 150–300 元，售后市场规模约 40 亿元/年')

doc.add_heading('二、行业痛点', level=2)
table = doc.add_table(rows=6, cols=3)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = '痛点'
hdr[1].text = '现状'
hdr[2].text = '影响'
data = [
    ('报修渠道分散', '400电话、微信、公众号多通道', '用户平均等待15分钟才能提交工单'),
    ('人工分诊准确率低', '客服凭经验判断', '30%工单派错工程师，二次上门'),
    ('派单效率低', '主管人工调度', '高峰时段派单延迟2小时以上'),
    ('备件携带不准', '工程师不知道带什么', '约18%工单现场发现缺件'),
    ('满意度无人跟进', '修完就结束', '低分投诉发现晚，口碑流失'),
]
for i, (a,b,c) in enumerate(data, 1):
    row = table.rows[i].cells
    row[0].text = a; row[1].text = b; row[2].text = c

doc.add_heading('三、竞争格局', level=2)
doc.add_paragraph('传统售后外包商：人力密集型，无AI能力，靠人堆')
doc.add_paragraph('品牌自建售后：重资产，只服务自家品牌')
doc.add_paragraph('新兴AI售后平台：本项目定位，用AI降本增效')

doc.add_heading('四、数智化趋势', level=2)
doc.add_paragraph('1. AI 意图识别：用户说一句话，自动归类故障类型和紧急程度')
doc.add_paragraph('2. 知识库检索：历史工单 + 维修手册向量检索，给出 Top3 可能原因')
doc.add_paragraph('3. 智能派单：距离 + 技能 + 负载三维度算法匹配')
doc.add_paragraph('4. 预测性维护：根据设备数据提前预警，主动上门')

doc.add_heading('五、目标客户', level=2)
doc.add_paragraph('中小型智能家居品牌（年出货 10–100 万台）')
doc.add_paragraph('区域型售后外包服务商')
doc.add_paragraph('物业增值服务部门')

doc.save(r'C:\Users\Lenovo\Desktop\f24015431王绍轩2\智家速修-AI售后工单\docs\行业分析.docx')
print('docx saved')
